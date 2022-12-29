import openpyxl
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

path = "E:/Python Projects/onizanalyzer/replays/txtfile output/#Humandata.xlsx"

hwb = openpyxl.load_workbook(filename=path)
hsheet = hwb.worksheets[0]

htab = Table(displayName="Table1", ref="B1:AZ79")
hsheet.add_table(htab)

print(hsheet)

